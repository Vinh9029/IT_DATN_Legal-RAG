"""
Nạp corpus từ nguồn dữ liệu CỤC BỘ — bất kể định dạng file.

─── Vì sao có file này ───

Nguồn HuggingFace `th1nhng0/vietnamese-legal-documents` chỉ là chỗ dựa tạm.
Nhóm đã chốt đi tìm nguồn luật dân sự riêng, và nguồn đó có thể về dưới bất kỳ
dạng nào: `.docx` tải từ trang tra cứu luật, `.csv` xuất từ Excel, `.jsonl` do
người khác crawl, `.parquet` từ một dataset khác. Nếu `corpus_filter` cứ gọi
thẳng `load_dataset()` thì mỗi lần đổi nguồn là một lần phải sửa code lõi.

Nên tách đôi trách nhiệm:
  - `corpus_loader` (file này) : ĐỌC gì cũng được → trả doc dict chuẩn hoá.
  - `corpus_filter`            : LỌC phạm vi dân sự trên doc dict chuẩn hoá đó.

Thêm dataset mới = thả file vào `data/corpus_civil/`. Không phải sửa dòng code
nào. Nếu tên cột lạ thì khai thêm bí danh vào `FIELD_ALIASES`
(`config/qa_settings.py`) — vẫn không phải đụng vào logic.

─── Hợp đồng dữ liệu (doc dict chuẩn hoá) ───

    {
      "id":             str,   # định danh trong nguồn, hoặc sinh từ nội dung
      "content":        str,   # plain text, đã bỏ HTML
      "content_length": int,
      "metadata":       {title, so_hieu, loai_van_ban, nganh, linh_vuc, ...},
      "source_doc_id":  str,   # khoá chống leakage lúc group split
      "source_file":    str,   # truy ngược khi phát hiện dữ liệu bẩn
    }

Đây đúng là hình dạng mà `pair_generator.generate_pair()` đang nhận, nên hai
đường nạp (cục bộ / HuggingFace) dùng chung toàn bộ phần phía sau.
"""

import csv
import json
import re
import shutil
import subprocess
import tempfile
import unicodedata
from collections import Counter
from pathlib import Path

from loguru import logger

from config.qa_settings import (
    ARTICLE_MIN_LENGTH,
    CONVERTED_DOC_DIR,
    FIELD_ALIASES,
    MIN_CONTENT_LENGTH,
    PDF_HEADER_REPEAT_RATIO,
    PDF_MIN_CHARS_PER_PAGE,
    PDF_REFLOW,
    SUPPORTED_SUFFIXES,
)


# ══════════════════════════════════════════════════════════════════
# Chuẩn hoá tên cột
# ══════════════════════════════════════════════════════════════════

def _slug(text) -> str:
    """
    Đưa tên cột về dạng so khớp được: bỏ dấu, thường hoá, gạch dưới.

    Nguồn dữ liệu người Việt xuất ra hay có header tiếng Việt CÓ DẤU
    ("Nội dung", "Số ký hiệu", "Loại văn bản"). So khớp thẳng chuỗi thì cứ đổi
    nguồn là trượt hết. Bỏ dấu trước rồi mới khớp: "Nội dung" thành `noi_dung`,
    trùng luôn với bí danh đã khai.
    """
    text = str(text or "").strip().lower()
    text = unicodedata.normalize("NFD", text)
    text = "".join(c for c in text if unicodedata.category(c) != "Mn")
    text = text.replace("đ", "d")
    text = re.sub(r"[^a-z0-9]+", "_", text)
    return text.strip("_")


def _slug_record(raw: dict) -> dict:
    """Bản sao của record với key đã slug hoá. Giữ nguyên giá trị."""
    return {_slug(k): v for k, v in raw.items()}


# Nhiều nguồn ghi ô trống thành chuỗi chứ không phải null
_EMPTY_STRINGS = {"none", "nan", "null", "n/a", "-", "chưa xác định"}


def _pick(slugged: dict, field: str) -> str:
    """Lấy giá trị đầu tiên không rỗng theo danh sách bí danh của `field`."""
    for alias in FIELD_ALIASES.get(field, []):
        value = slugged.get(_slug(alias))
        if value is None:
            continue
        text = str(value).strip()
        if text and text.lower() not in _EMPTY_STRINGS:
            return text
    return ""


# ══════════════════════════════════════════════════════════════════
# Đọc file thô thành list record dict
# ══════════════════════════════════════════════════════════════════

_HTML_RE = re.compile(r"<\s*(p|div|br|span|table|tr|td|h[1-6]|body|html)\b", re.IGNORECASE)


def _read_text(path: Path) -> str:
    """
    Đọc file text, thử lần lượt vài bảng mã.

    `utf-8-sig` đứng đầu vì Excel trên Windows xuất CSV kèm BOM; không xử lý
    BOM thì cột đầu tiên mang tên rác và mọi bí danh đều trượt.
    """
    for encoding in ("utf-8-sig", "utf-8", "cp1258", "cp1252", "latin-1"):
        try:
            return path.read_text(encoding=encoding)
        except UnicodeDecodeError:
            continue
    return path.read_bytes().decode("utf-8", errors="replace")


def _read_jsonl(path: Path) -> list[dict]:
    records = []
    for line_no, line in enumerate(_read_text(path).splitlines(), 1):
        line = line.strip()
        if not line:
            continue
        try:
            obj = json.loads(line)
        except json.JSONDecodeError as exc:
            logger.warning(f"{path.name}:{line_no} — bỏ qua dòng JSON hỏng ({exc})")
            continue
        if isinstance(obj, dict):
            records.append(obj)
    return records


def _read_json(path: Path) -> list[dict]:
    data = json.loads(_read_text(path))
    if isinstance(data, dict):
        # Bọc kiểu {"data": [...]} / {"documents": [...]} rất phổ biến
        for key in ("data", "documents", "docs", "records", "items", "results"):
            if isinstance(data.get(key), list):
                data = data[key]
                break
        else:
            data = [data]
    return [d for d in data if isinstance(d, dict)]


def _read_csv(path: Path) -> list[dict]:
    text = _read_text(path)
    try:
        dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|")
    except csv.Error:
        # Sniffer hay thua khi ô chứa văn bản luật dài có đủ loại dấu câu.
        # Đoán theo phần mở rộng thay vì bỏ cuộc.
        dialect = csv.excel_tab if path.suffix.lower() == ".tsv" else csv.excel
    # Văn bản luật dễ vượt giới hạn field mặc định (128KB) của module csv
    csv.field_size_limit(2 ** 31 - 1)
    return [dict(row) for row in csv.DictReader(text.splitlines(), dialect=dialect)]


def _read_xlsx(path: Path) -> list[dict]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ImportError(
            f"Cần `openpyxl` để đọc {path.name}. Cài: pip install openpyxl"
        ) from exc

    workbook = load_workbook(path, read_only=True, data_only=True)
    records = []
    for sheet in workbook.worksheets:
        rows = sheet.iter_rows(values_only=True)
        try:
            header = [str(c) if c is not None else "" for c in next(rows)]
        except StopIteration:
            continue
        for row in rows:
            if all(c is None or str(c).strip() == "" for c in row):
                continue
            records.append({h: v for h, v in zip(header, row) if h})
    workbook.close()
    return records


def _read_docx(path: Path) -> list[dict]:
    """
    Một file .docx là MỘT văn bản luật.

    Lấy cả text trong bảng: biểu mẫu, phụ lục, bảng mức bồi thường trong văn
    bản luật Việt Nam thường nằm trong table, bỏ qua là mất đúng phần nội dung
    đáng hỏi nhất.
    """
    try:
        import docx
    except ImportError as exc:
        raise ImportError(
            f"Cần `python-docx` để đọc {path.name}. Cài: pip install python-docx"
        ) from exc

    document = docx.Document(str(path))
    parts = [p.text for p in document.paragraphs if p.text and p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text and c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))

    return [{"title": path.stem, "content": "\n".join(parts)}]


def _read_plain(path: Path) -> list[dict]:
    """Một file .txt/.md/.html là MỘT văn bản luật; tên file làm tiêu đề."""
    return [{"title": path.stem, "content": _read_text(path)}]


# ══════════════════════════════════════════════════════════════════
# PDF
# ══════════════════════════════════════════════════════════════════
# PDF không lưu văn bản, nó lưu VỊ TRÍ CỦA TỪNG KÝ TỰ trên từng trang. Nên
# ngoài việc trích chữ ra, còn phải dọn hai thứ mà bố cục theo trang để lại:
# header/footer chạy suốt văn bản, và câu bị cắt ngang ở cuối mỗi dòng.


def _pdf_pages(path: Path) -> list[str]:
    """
    Trả về text của từng trang. Thử lần lượt 3 thư viện.

    Xếp theo chất lượng trích văn bản tiếng Việt, không theo độ phổ biến:
    pymupdf giữ đúng thứ tự đọc và dấu thanh; pdfplumber chậm hơn nhưng xử lý
    bảng tốt; pypdf là chốt chặn cuối vì thuần Python, ở đâu cũng cài được.
    """
    errors = []

    try:
        try:
            import pymupdf
        except ImportError:
            import fitz as pymupdf
        with pymupdf.open(str(path)) as document:
            return [page.get_text("text") or "" for page in document]
    except ImportError as exc:
        errors.append(f"pymupdf: {exc}")

    try:
        import pdfplumber
        with pdfplumber.open(str(path)) as pdf:
            return [(page.extract_text() or "") for page in pdf.pages]
    except ImportError as exc:
        errors.append(f"pdfplumber: {exc}")

    try:
        from pypdf import PdfReader
        return [(page.extract_text() or "") for page in PdfReader(str(path)).pages]
    except ImportError as exc:
        errors.append(f"pypdf: {exc}")

    raise ImportError(
        f"Cần một thư viện đọc PDF để đọc {path.name}. "
        f"Cài một trong: pip install pymupdf  (khuyến nghị) / pdfplumber / pypdf. "
        f"Chi tiết: {'; '.join(errors)}"
    )


# "12", "- 3 -", "Trang 4/120", "Page 2" — đứng một mình trên một dòng
_PDF_PAGE_NUMBER_RE = re.compile(
    r"^(?:[-–—\s]*\d{1,4}[-–—\s]*|(?:trang|page)\s*\d{1,4}(?:\s*/\s*\d{1,4})?)$",
    re.IGNORECASE,
)

# Dòng MỞ ĐẦU một đơn vị cấu trúc — không bao giờ được nối vào dòng trước.
# "Điều N" CỐ Ý không nằm trong danh sách này: nó được xét riêng bằng
# `_looks_like_article_heading` để trích dẫn chéo ("Điều 53 của Bộ luật này")
# được nối trả về đoạn văn của nó, thay vì đứng riêng một dòng rồi bị nhầm
# thành tiêu đề Điều ở bước cắt.
_PDF_STRUCTURE_RE = re.compile(
    r"^(?:Chương\s|Mục\s|Phần\s|Tiểu\s*mục|PHỤ\s*LỤC|Căn\s+cứ\b"
    r"|\d+[.)]\s|[a-zđ][.)]\s|\(\d+\)|[-–—•*]\s)",
    re.IGNORECASE,
)


def _starts_new_block(line: str) -> bool:
    """Dòng này mở đầu một đơn vị cấu trúc mới (nên KHÔNG nối vào dòng trước)?"""
    if _PDF_STRUCTURE_RE.match(line):
        return True
    # `_ARTICLE_RE` khai ở cuối file (mục "Cắt theo Điều"); Python tra tên lúc
    # gọi nên dùng trước chỗ khai được, và giữ được MỘT chỗ định nghĩa duy nhất
    # thế nào là tiêu đề Điều.
    match = _ARTICLE_RE.match(line)
    return bool(match) and _looks_like_article_heading(match, line)

_PDF_SENTENCE_END_RE = re.compile(r"[.:;!?…][\"'”’)\]]?$")

# Dòng ngắn hơn ngưỡng này gần như chắc chắn là dòng KẾT một đoạn (hoặc một
# tiêu đề), không phải câu bị cắt vì hết chiều ngang trang → không nối.
_PDF_REFLOW_MIN_LINE = 40


def _reflow_pdf_text(text: str) -> str:
    """
    Nối lại những dòng bị PDF cắt ngang giữa câu.

    Chỉ nối khi CẢ BA điều kiện cùng đúng: dòng trước đủ dài để tin là bị cắt
    vì hết chiều ngang trang, không kết thúc bằng dấu câu, và dòng sau không mở
    đầu một đơn vị cấu trúc mới. Ba điều kiện là để thà bỏ sót còn hơn nối
    nhầm — nối nhầm "Điều 6" vào cuối Điều 5 thì `split_into_articles` mất luôn
    một Điều, mà mất thì không có triệu chứng nào nhìn thấy được.
    """
    lines: list[str] = []
    for raw in text.split("\n"):
        current = raw.strip()
        previous = lines[-1] if lines else ""
        if (current and previous
                and len(previous) >= _PDF_REFLOW_MIN_LINE
                and not previous.isupper()
                and not _PDF_SENTENCE_END_RE.search(previous)
                and not _starts_new_block(current)):
            lines[-1] = f"{previous} {current}"
        else:
            lines.append(current)
    return "\n".join(lines)


def _clean_pdf_pages(pages: list[str]) -> str:
    """Ghép các trang thành một văn bản, bỏ header/footer và số trang."""
    page_lines = [[line.strip() for line in (page or "").splitlines()] for page in pages]

    # Header/footer nhận diện bằng TẦN SUẤT chứ không bằng vị trí: "CÔNG BÁO/Số
    # 45", tên cơ quan, số hiệu văn bản in lại ở mọi trang. Vị trí không dùng
    # được vì mỗi nơi phát hành đặt một kiểu.
    repeated: set[str] = set()
    if len(pages) >= 3:
        counter: Counter = Counter()
        for lines in page_lines:
            for line in {ln for ln in lines if ln and len(ln) <= 80}:
                counter[line] += 1
        limit = max(2, int(len(pages) * PDF_HEADER_REPEAT_RATIO))
        repeated = {line for line, count in counter.items() if count >= limit}

    kept: list[str] = []
    for lines in page_lines:
        for line in lines:
            if not line:
                kept.append("")
            elif line in repeated or _PDF_PAGE_NUMBER_RE.match(line):
                continue
            else:
                kept.append(line)
        kept.append("")  # ranh giới trang

    text = "\n".join(kept)
    if PDF_REFLOW:
        text = _reflow_pdf_text(text)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _read_pdf(path: Path) -> list[dict]:
    """Một file .pdf là MỘT văn bản luật."""
    pages = _pdf_pages(path)
    content = _clean_pdf_pages(pages)

    # PDF scan (ảnh chụp) trích ra gần như 0 ký tự. Không báo ở đây thì nó lặng
    # lẽ bị loại ở bước "content quá ngắn" và người dùng chỉ thấy corpus hụt
    # mất vài file mà không hiểu vì sao.
    if pages and len(content) / len(pages) < PDF_MIN_CHARS_PER_PAGE:
        logger.warning(
            f"{path.name}: chỉ trích được {len(content) / len(pages):.0f} ký tự/trang "
            f"trên {len(pages)} trang → gần như chắc chắn là PDF SCAN (ảnh), cần OCR. "
            f"Cách xử lý: chạy OCRmyPDF (`ocrmypdf -l vie in.pdf out.pdf`) rồi thả bản "
            f"đã OCR vào, hoặc tìm bản .docx/.html của cùng văn bản đó."
        )

    return [{"title": path.stem, "content": content}]


# ══════════════════════════════════════════════════════════════════
# RTF
# ══════════════════════════════════════════════════════════════════

# Các "destination" của RTF chứa dữ liệu điều khiển, không phải nội dung
_RTF_DESTINATIONS = frozenset((
    "aftncn", "aftnsep", "aftnsepc", "annotation", "atnauthor", "atndate", "atnicn",
    "atnid", "atnparent", "atnref", "atntime", "atrfend", "atrfstart", "author",
    "background", "bkmkend", "bkmkstart", "blipuid", "buptim", "category",
    "colorschememapping", "colortbl", "comment", "company", "creatim", "datafield",
    "datastore", "defchp", "defpap", "do", "doccomm", "docvar", "dptxbxtext",
    "falt", "fchars", "ffdeftext", "ffentrymcr", "ffexitmcr", "ffformat", "ffhelptext",
    "ffl", "ffname", "ffstattext", "field", "file", "filetbl", "fldinst", "fldtype",
    "fname", "fontemb", "fontfile", "fonttbl", "footer", "footerf", "footerl",
    "footerr", "footnote", "formfield", "ftncn", "ftnsep", "ftnsepc", "g",
    "generator", "gridtbl", "header", "headerf", "headerl", "headerr", "hl", "hlfr",
    "hlinkbase", "hlloc", "hlsrc", "hsv", "htmltag", "info", "keycode", "keywords",
    "latentstyles", "lchars", "levelnumbers", "leveltext", "lfolevel", "linkval",
    "list", "listlevel", "listname", "listoverride", "listoverridetable",
    "listpicture", "liststylename", "listtable", "listtext", "lsdlockedexcept",
    "mailmerge", "manager", "mhtmltag", "nesttableprops", "nextfile", "nonesttables",
    "objalias", "objclass", "objdata", "object", "objname", "objsect", "objtime",
    "oldcprops", "oldpprops", "oldsprops", "oldtprops", "oleclsid", "operator",
    "panose", "password", "passwordhash", "pgp", "pgptbl", "picprop", "pict", "pn",
    "pnseclvl", "pntext", "pntxta", "pntxtb", "printim", "private", "propname",
    "protend", "protstart", "protusertbl", "pxe", "result", "revtbl", "revtim",
    "rsidtbl", "rxe", "shp", "shpgrp", "shpinst", "shppict", "shprslt", "shptxt",
    "sn", "sp", "staticval", "stylesheet", "subject", "sv", "svb", "tc", "template",
    "themedata", "title", "txe", "ud", "upr", "userprops", "wgrffmtfilter",
    "windowcaption", "writereservation", "writereservhash", "xe", "xform",
    "xmlattrname", "xmlattrvalue", "xmlclose", "xmlname", "xmlnstbl", "xmlopen",
))

_RTF_SPECIAL = {
    "par": "\n", "sect": "\n\n", "page": "\n\n", "line": "\n", "tab": "\t",
    "emdash": "\u2014", "endash": "\u2013", "emspace": "\u2003", "enspace": "\u2002",
    "qmspace": "\u200a", "bullet": "\u2022", "lquote": "\u2018", "rquote": "\u2019",
    "ldblquote": "\u201c", "rdblquote": "\u201d", "row": "\n", "cell": " | ",
    "nestcell": " | ", "nestrow": "\n",
}

_RTF_TOKEN_RE = re.compile(
    r"\\([a-z]{1,32})(-?\d{1,10})?[ ]?|\\'([0-9a-f]{2})|\\([^a-z])|([{}])|[\r\n]+|(.)",
    re.IGNORECASE | re.DOTALL,
)


def _rtf_to_text(source: str) -> str:
    """
    Bóc text khỏi RTF mà không cần thư viện ngoài.

    Vì sao đáng tự viết: `.doc` tải từ các trang tra cứu luật RẤT hay thực chất
    là RTF đổi đuôi (xem `_sniff_binary_kind`), nên đây là đường đi thường gặp
    chứ không phải trường hợp hiếm.

    Điểm phải làm đúng cho tiếng Việt: Word ghi ký tự có dấu thành `\\uNNNN`
    KÈM ký tự thay thế phía sau cho trình đọc đời cũ. Không bỏ đúng số ký tự
    thay thế đó (`\\ucN` khai báo) thì mỗi chữ có dấu kéo theo một ký tự rác.
    """
    codepage = "cp1252"
    match = re.search(r"\\ansicpg(\d{3,5})", source[:2048])
    if match:
        codepage = f"cp{match.group(1)}"

    stack: list[tuple[int, bool]] = []
    ignorable = False
    ucskip = 1      # số ký tự thay thế phải bỏ sau mỗi \uNNNN
    curskip = 0
    out: list[str] = []

    for token in _RTF_TOKEN_RE.finditer(source):
        word, arg, hexcode, char, brace, text = token.groups()

        if brace:
            curskip = 0
            if brace == "{":
                stack.append((ucskip, ignorable))
            elif stack:
                ucskip, ignorable = stack.pop()
        elif char:
            curskip = 0
            if char == "~":
                if not ignorable:
                    out.append("\xa0")
            elif char in "{}\\":
                if not ignorable:
                    out.append(char)
            elif char == "*":
                ignorable = True
        elif word:
            curskip = 0
            if word in _RTF_DESTINATIONS:
                ignorable = True
            elif ignorable:
                pass
            elif word in _RTF_SPECIAL:
                out.append(_RTF_SPECIAL[word])
            elif word == "uc":
                ucskip = int(arg or 1)
            elif word == "u":
                code = int(arg or 0)
                if code < 0:
                    code += 0x10000
                out.append(chr(code) if 0 <= code <= 0x10FFFF else "?")
                curskip = ucskip
        elif hexcode:
            if curskip > 0:
                curskip -= 1
            elif not ignorable:
                byte = bytes([int(hexcode, 16)])
                try:
                    out.append(byte.decode(codepage, errors="replace"))
                except LookupError:
                    out.append(byte.decode("cp1252", errors="replace"))
        elif text:
            if curskip > 0:
                curskip -= 1
            elif not ignorable:
                out.append(text)

    result = "".join(out)
    result = re.sub(r"[ \t]+", " ", result)
    return re.sub(r"\n{3,}", "\n\n", result).strip()


def _read_rtf(path: Path) -> list[dict]:
    """Một file .rtf là MỘT văn bản luật."""
    # `striprtf` xử lý bảng và font table kỹ hơn; dùng nếu có, không thì tự bóc.
    raw = _read_text(path)
    try:
        from striprtf.striprtf import rtf_to_text
        content = rtf_to_text(raw, errors="ignore")
    except ImportError:
        content = _rtf_to_text(raw)
    return [{"title": path.stem, "content": content.strip()}]


# ══════════════════════════════════════════════════════════════════
# .doc (Word nhị phân đời cũ)
# ══════════════════════════════════════════════════════════════════

def _sniff_binary_kind(path: Path) -> str:
    """
    Đuôi file KHÔNG nói lên định dạng thật. Đọc mấy byte đầu để biết chắc.

    Đây không phải đề phòng suông: các trang tra cứu luật Việt Nam thường xuất
    "bản Word" bằng cách đổi đuôi một file HTML hoặc RTF thành `.doc`. Tin vào
    đuôi file thì `python-docx` ném lỗi "not a zip file" và cả văn bản bị bỏ.
    """
    with open(path, "rb") as f:
        head = f.read(512)

    if head[:4] == b"PK\x03\x04":
        return "docx"                                   # thực chất là .docx
    if head[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
        return "ole"                                    # Word 97-2003 thật
    stripped = head.lstrip(b"\xef\xbb\xbf \t\r\n")
    if stripped[:5].lower() == b"{\\rtf":
        return "rtf"
    lowered = stripped[:256].lower()
    if lowered[:1] == b"<" or b"<html" in lowered or b"<?xml" in lowered:
        return "html"
    return "unknown"


def _find_soffice() -> str:
    """Tìm LibreOffice: PATH trước, rồi các vị trí cài mặc định."""
    for name in ("soffice", "soffice.exe", "libreoffice"):
        found = shutil.which(name)
        if found:
            return found
    for candidate in (
        r"C:\Program Files\LibreOffice\program\soffice.exe",
        r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
        "/usr/bin/soffice", "/usr/bin/libreoffice",
        "/Applications/LibreOffice.app/Contents/MacOS/soffice",
    ):
        if Path(candidate).exists():
            return candidate
    return ""


def _convert_via_soffice(path: Path, target: Path) -> bool:
    exe = _find_soffice()
    if not exe:
        return False

    # Convert vào thư mục tạm rồi mới đổi tên: soffice tự đặt tên đầu ra theo
    # tên file nguồn, không cho chỉ định.
    with tempfile.TemporaryDirectory() as tmp:
        result = subprocess.run(
            [exe, "--headless", "--norestore", "--convert-to", "docx",
             "--outdir", tmp, str(path.resolve())],
            capture_output=True, timeout=300,
        )
        produced = Path(tmp) / f"{path.stem}.docx"
        if not produced.exists():
            logger.debug(f"soffice không tạo được {produced.name}: {result.stderr[:200]!r}")
            return False
        shutil.move(str(produced), str(target))
    return True


def _convert_via_word(path: Path, target: Path) -> bool:
    """Nhờ Microsoft Word convert (chỉ Windows, cần đã cài Word + pywin32)."""
    try:
        import pythoncom
        import win32com.client
    except ImportError:
        return False

    pythoncom.CoInitialize()
    word = None
    document = None
    try:
        word = win32com.client.Dispatch("Word.Application")
        word.Visible = False
        word.DisplayAlerts = False
        document = word.Documents.Open(
            str(path.resolve()), ReadOnly=True, AddToRecentFiles=False, Visible=False
        )
        document.SaveAs2(str(target.resolve()), FileFormat=16)  # 16 = wdFormatDocumentDefault
        return True
    finally:
        try:
            if document is not None:
                document.Close(False)
        except Exception:
            pass
        try:
            if word is not None:
                word.Quit()
        except Exception:
            pass
        pythoncom.CoUninitialize()


def _convert_doc_to_docx(path: Path) -> Path:
    """
    Convert `.doc` nhị phân sang `.docx` rồi trả về đường dẫn bản đã convert.

    Không có thư viện Python thuần nào đọc `.doc` nhị phân mà ra tiếng Việt
    đúng dấu — định dạng đó lưu text qua piece table nén, không phải một khối
    liền. Nên mượn tay công cụ đã làm đúng việc: LibreOffice hoặc Word.

    Bản convert được CACHE theo (đường dẫn, kích thước, mtime): mỗi lần convert
    mất vài giây, mà `--rebuild-corpus` thì chạy lại rất nhiều lần.
    """
    CONVERTED_DOC_DIR.mkdir(parents=True, exist_ok=True)

    from evol_instruct.src.utils import generate_item_id
    stat = path.stat()
    key = generate_item_id(f"{path.resolve()}:{stat.st_size}:{int(stat.st_mtime)}")
    target = CONVERTED_DOC_DIR / f"{path.stem}__{key}.docx"
    if target.exists():
        logger.debug(f"  {path.name}: dùng bản convert đã cache")
        return target

    for name, convert in (("LibreOffice", _convert_via_soffice), ("MS Word", _convert_via_word)):
        try:
            if convert(path, target) and target.exists():
                logger.info(f"  {path.name} → .docx qua {name}")
                return target
        except Exception as exc:
            logger.debug(f"  {path.name}: {name} thất bại — {type(exc).__name__}: {exc}")

    raise RuntimeError(
        f"Không convert được {path.name} (.doc nhị phân Word 97-2003). Chọn một cách:\n"
        f"  1. Cài LibreOffice (https://www.libreoffice.org) — convert được hàng loạt, "
        f"không phải thao tác tay;\n"
        f"  2. Đã có Microsoft Word thì: pip install pywin32;\n"
        f"  3. Mở file bằng Word rồi Save As → .docx, thả bản .docx vào thư mục corpus."
    )


def _read_doc(path: Path) -> list[dict]:
    """Một file .doc là MỘT văn bản luật — nhưng bên trong có thể là 4 định dạng."""
    kind = _sniff_binary_kind(path)

    if kind == "docx":
        logger.info(f"  {path.name}: đuôi .doc nhưng ruột là .docx")
        return _read_docx(path)
    if kind == "rtf":
        logger.info(f"  {path.name}: đuôi .doc nhưng ruột là RTF")
        return _read_rtf(path)
    if kind == "html":
        logger.info(f"  {path.name}: đuôi .doc nhưng ruột là HTML")
        return _read_plain(path)

    if kind != "ole":
        logger.warning(f"{path.name}: không nhận ra định dạng, thử coi như Word nhị phân")

    records = _read_docx(_convert_doc_to_docx(path))
    for record in records:
        # Bản convert mang tên đã gắn hash cache; tiêu đề phải là tên file gốc
        record["title"] = path.stem
    return records


def _read_parquet(path: Path) -> list[dict]:
    try:
        import pyarrow.parquet as pq
    except ImportError as exc:
        raise ImportError(
            f"Cần `pyarrow` để đọc {path.name}. Cài: pip install pyarrow"
        ) from exc
    return pq.read_table(path).to_pylist()


_READERS = {
    ".jsonl": _read_jsonl, ".ndjson": _read_jsonl,
    ".json": _read_json,
    ".csv": _read_csv, ".tsv": _read_csv,
    ".xlsx": _read_xlsx, ".xlsm": _read_xlsx,
    ".docx": _read_docx,
    ".doc": _read_doc,
    ".rtf": _read_rtf,
    ".pdf": _read_pdf,
    ".txt": _read_plain, ".md": _read_plain,
    ".html": _read_plain, ".htm": _read_plain,
    ".parquet": _read_parquet,
}

# Hai nơi khai định dạng phải khớp nhau, nếu không thông báo lỗi cho người
# dùng sẽ liệt kê một đằng còn code đọc được một nẻo.
assert set(_READERS) == set(SUPPORTED_SUFFIXES), (
    "SUPPORTED_SUFFIXES trong config/qa_settings.py đã lệch với _READERS. Sửa cả hai."
)


# ══════════════════════════════════════════════════════════════════
# Chuẩn hoá record thành doc dict
# ══════════════════════════════════════════════════════════════════

def _to_plain_text(value) -> str:
    text = "" if value is None else str(value)
    if not _HTML_RE.search(text):
        return text.strip()

    # Import muộn: `src.data_loader` kéo theo `datasets` và `config.settings`.
    # Đường chạy cục bộ không cần hai thứ đó và không nên chết vì thiếu chúng.
    try:
        from evol_instruct.src.data_loader import clean_html
        return clean_html(text)
    except Exception:
        from bs4 import BeautifulSoup
        return BeautifulSoup(text, "html.parser").get_text(separator="\n").strip()


# Các trường metadata bê nguyên nếu nguồn có
_METADATA_FIELDS = (
    "title", "so_hieu", "loai_van_ban", "nganh", "linh_vuc",
    "co_quan_ban_hanh", "ngay_ban_hanh", "tinh_trang_hieu_luc",
)


def normalize_record(raw: dict, source_file: str = "", index: int = 0) -> dict | None:
    """
    Ép một record thô bất kỳ về doc dict chuẩn hoá.

    Trả None khi record không có nội dung — đó là chuyện bình thường (dòng
    trống trong CSV, sheet phụ trong Excel), không phải lỗi cần báo.
    """
    slugged = _slug_record(raw)
    content = _to_plain_text(_pick(slugged, "content"))
    if not content:
        return None

    metadata = {field: _pick(slugged, field) for field in _METADATA_FIELDS}
    # Giữ lại mọi cột lạ của nguồn: khi phát hiện dữ liệu bẩn còn có cái mà
    # truy, và nguồn mới có thể mang trường hữu ích mà ta chưa biết trước.
    for key, value in slugged.items():
        if key not in metadata and value not in (None, ""):
            metadata.setdefault(key, str(value).strip())

    doc_id = _pick(slugged, "id")
    if not doc_id:
        from evol_instruct.src.utils import generate_item_id
        seed = metadata.get("so_hieu") or metadata.get("title") or content[:500]
        doc_id = generate_item_id(f"{source_file}:{index}:{seed}")

    return {
        "id": str(doc_id),
        "content": content,
        "content_length": len(content),
        "metadata": metadata,
        "source_doc_id": f"doc_{doc_id}",
        "source_file": source_file,
    }


# ══════════════════════════════════════════════════════════════════
# Điểm vào
# ══════════════════════════════════════════════════════════════════

def iter_source_files(source) -> list[Path]:
    """Liệt kê file đọc được trong `source` (file đơn hoặc thư mục, đệ quy)."""
    source = Path(source)
    if source.is_file():
        return [source] if source.suffix.lower() in _READERS else []
    if not source.is_dir():
        return []
    return sorted(
        p for p in source.rglob("*")
        # "~$..." là file khoá tạm của Word/Excel, mở ra là hỏng
        if p.is_file() and p.suffix.lower() in _READERS and not p.name.startswith("~$")
        # Bỏ thư mục ẩn: chỗ để bản nháp, bản backup, hoặc bản convert tạm —
        # quét vào là cùng một văn bản vào corpus hai lần.
        and not any(part.startswith(".") for part in p.relative_to(source).parts[:-1])
        # `data/corpus_civil/README.md` là tài liệu hướng dẫn nằm SẴN trong thư
        # mục corpus và có đuôi đọc được, nên nó tự vào corpus thành một
        # "văn bản luật" — đã xảy ra thật, ra một document `scope=unknown`.
        and p.stem.lower() != "readme"
    )


def has_local_corpus(source) -> bool:
    """Có file nào đọc được ở `source` không? Quyết định chọn đường nạp."""
    return bool(source) and bool(iter_source_files(source))


def load_local_corpus(source, min_content_length: int | None = None) -> list[dict]:
    """
    Nạp toàn bộ văn bản từ `source` về danh sách doc dict chuẩn hoá.

    Args:
        source: file hoặc thư mục. Thư mục thì quét đệ quy.
        min_content_length: ngưỡng độ dài; None thì lấy `MIN_CONTENT_LENGTH`.

    Trùng `source_doc_id` bị loại (giữ bản đầu): nguồn tải tay rất hay có file
    lặp, mà trùng document giữa train và test là data leakage kinh điển.
    """
    threshold = MIN_CONTENT_LENGTH if min_content_length is None else min_content_length
    files = iter_source_files(source)
    if not files:
        logger.warning(f"Không tìm thấy file đọc được nào trong: {source}")
        return []

    logger.info(f"Nạp corpus cục bộ từ {source} — {len(files)} file")

    documents: list[dict] = []
    seen: set[str] = set()
    skipped_short = 0
    skipped_dup = 0

    for path in files:
        reader = _READERS[path.suffix.lower()]
        try:
            raw_records = reader(path)
        except Exception as exc:
            # Một file hỏng không được phép làm hỏng cả mẻ nạp
            logger.error(f"Bỏ qua {path.name}: {type(exc).__name__}: {exc}")
            continue

        n_before = len(documents)
        for index, raw in enumerate(raw_records):
            doc = normalize_record(raw, source_file=path.name, index=index)
            if doc is None:
                continue
            if doc["content_length"] < threshold:
                skipped_short += 1
                continue
            if doc["source_doc_id"] in seen:
                skipped_dup += 1
                continue
            seen.add(doc["source_doc_id"])
            documents.append(doc)

        logger.debug(f"  {path.name}: {len(documents) - n_before} document")

    logger.info(
        f"Corpus cục bộ: {len(documents)} document "
        f"({skipped_short} bị loại vì content < {threshold} ký tự, "
        f"{skipped_dup} trùng source_doc_id)"
    )
    return documents


# ══════════════════════════════════════════════════════════════════
# Cắt theo Điều (tuỳ chọn — xem QA_SPLIT_BY_ARTICLE)
# ══════════════════════════════════════════════════════════════════

# "Điều 5.", "Điều 12:", "Điều 468a." — đầu dòng, có số, có dấu ngắt.
# Cố ý KHÔNG bắt "điều" viết thường giữa câu ("theo điều 5 nói trên").
# Nhóm 2 là dấu ngắt, tách riêng vì `_looks_like_article_heading` cần biết
# có hay không — xem docstring hàm đó.
_ARTICLE_RE = re.compile(r"^[ \t]*Điều\s+(\d+[a-zđ]?)\s*([.:．])?(?:\s|$)", re.MULTILINE)


def _looks_like_article_heading(match: re.Match, text: str) -> bool:
    """
    Khớp được "Điều N" ở đầu dòng rồi, nhưng đó là TIÊU ĐỀ Điều hay là TRÍCH
    DẪN tới Điều khác?

    Vì sao phải phân biệt: văn bản luật trích dẫn chéo liên tục ("...quy định
    tại Điều 53 của Bộ luật này..."), và khi trích dẫn rơi đúng đầu dòng — rất
    hay gặp ở PDF vì chữ xuống dòng theo chiều ngang trang — thì nó khớp y hệt
    một tiêu đề. Đo trên Bộ luật Dân sự 2015: 10 chỗ, Bộ luật Tố tụng Dân sự
    2015: 21 chỗ. Mỗi chỗ hỏng KÉP: sinh ra một "Điều" ma bắt đầu từ giữa câu,
    đồng thời cắt cụt Điều thật đang chứa nó.

    Hai dấu hiệu của tiêu đề thật:
      - có dấu ngắt ngay sau số ("Điều 53." / "Điều 53:"), hoặc
      - không có dấu ngắt nhưng chữ kế tiếp VIẾT HOA — tên Điều luôn viết hoa
        ("Điều 53 Người giám hộ đương nhiên"), còn trích dẫn thì đi tiếp bằng
        chữ thường ("Điều 53 của Bộ luật này").
    """
    if match.group(2):
        return True
    tail = text[match.end():match.end() + 1]
    return bool(tail) and tail.isupper()


def split_into_articles(content: str, min_length: int | None = None) -> list[tuple[str, str]]:
    """
    Cắt nội dung văn bản thành từng Điều.

    Returns:
        list[(số điều, nội dung điều)]. Rỗng khi văn bản không đánh số theo
        Điều (Thông tư đánh `I.`/`1.`, văn bản sửa đổi...) — caller giữ nguyên
        cả văn bản trong trường hợp đó.
    """
    threshold = ARTICLE_MIN_LENGTH if min_length is None else min_length
    content = content or ""
    matches = [m for m in _ARTICLE_RE.finditer(content)
               if _looks_like_article_heading(m, content)]
    if len(matches) < 2:
        return []

    articles: list[tuple[str, str]] = []
    for i, match in enumerate(matches):
        end = matches[i + 1].start() if i + 1 < len(matches) else len(content)
        body = content[match.start():end].strip()
        if len(body) >= threshold:
            articles.append((match.group(1), body))
    return articles


def expand_documents_by_article(documents: list[dict]) -> list[dict]:
    """
    Biến mỗi văn bản thành nhiều "document con", mỗi Điều một cái.

    Vì sao cần: phạm vi dân sự là ÍT VĂN BẢN NHƯNG RẤT DÀI (riêng Bộ luật Dân
    sự 2015 có 689 Điều). Lấy 1 cặp câu hỏi trên mỗi văn bản thì cả bộ luật
    lớn nhất cũng chỉ đóng góp đúng 1 cặp, mà 2.000 ký tự đầu của nó thì toàn
    phần "Căn cứ..." — nội dung đáng hỏi nằm ở giữa và không bao giờ tới lượt.

    `source_doc_id` chuyển thành `<doc>#dieu-<n>` và ĐÓ MỚI là khoá group
    split. Cách đọc: hai câu của cùng một Điều không bao giờ tách sang hai
    split khác nhau (đây là leakage thật cần chặn), còn các Điều khác nhau
    trong cùng một bộ luật thì được phép nằm khác split — nếu không thì cả bộ
    luật rơi trọn vào một split và không chia nổi train/val/test. Phải ghi rõ
    cách đọc này ở phần Limitations của báo cáo.
    """
    expanded: list[dict] = []
    n_split = 0

    for doc in documents:
        articles = split_into_articles(doc.get("content", ""))
        if not articles:
            expanded.append(doc)
            continue

        n_split += 1
        for number, body in articles:
            child = dict(doc)
            child["content"] = body
            child["content_length"] = len(body)
            child["source_doc_id"] = f"{doc['source_doc_id']}#dieu-{number}"
            child["id"] = f"{doc['id']}#dieu-{number}"
            child["parent_doc_id"] = doc["source_doc_id"]
            child["article_number"] = number
            child["metadata"] = dict(doc.get("metadata") or {})
            child["metadata"]["dieu"] = number
            expanded.append(child)

    logger.info(
        f"Cắt theo Điều: {n_split}/{len(documents)} văn bản cắt được "
        f"→ {len(expanded)} document con"
    )
    return expanded
