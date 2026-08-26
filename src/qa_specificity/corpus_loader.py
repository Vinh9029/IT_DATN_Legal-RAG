"""
Nạp corpus từ nguồn dữ liệu CỤC BỘ — bất kể định dạng file.

─── Vì sao có file này ───

Nguồn HuggingFace `th1nhng0/vietnamese-legal-documents` chỉ là chỗ dựa tạm.
Nhóm đã chốt đi tìm nguồn luật dân sự riêng, và nguồn đó có thể về dưới bất kỳ
dạng nào: `.docx` tải từ trang tra cứu luật, `.csv` xuất từ Excel, `.jsonl` do
người khác crawl, `.parquet` từ một dataset khác. Nếu `corpus_filter` cứ gọi
thẳng `load_dataset()` thì mỗi lần đổi nguồn là một lần phải sửa code lõi.

Nên tách đôi trách nhiệm:
  - `corpus_loader` (file này) : ĐỌC gì cũng được → trả doc dict chuẩn hoá.
  - `corpus_filter`            : LỌC phạm vi dân sự trên doc dict chuẩn hoá đó.

Thêm dataset mới = thả file vào `data/corpus_civil/`. Không phải sửa dòng code
nào. Nếu tên cột lạ thì khai thêm bí danh vào `FIELD_ALIASES`
(`config/qa_settings.py`) — vẫn không phải đụng vào logic.

─── Hợp đồng dữ liệu (doc dict chuẩn hoá) ───

    {
      "id":             str,   # định danh trong nguồn, hoặc sinh từ nội dung
      "content":        str,   # plain text, đã bỏ HTML
      "content_length": int,
      "metadata":       {title, so_hieu, loai_van_ban, nganh, linh_vuc, ...},
      "source_doc_id":  str,   # khoá chống leakage lúc group split
      "source_file":    str,   # truy ngược khi phát hiện dữ liệu bẩn
    }

Đây đúng là hình dạng mà `pair_generator.generate_pair()` đang nhận, nên hai
đường nạp (cục bộ / HuggingFace) dùng chung toàn bộ phần phía sau.
"""

import csv
import json
import re
import unicodedata
from pathlib import Path

from loguru import logger

from config.qa_settings import (
    ARTICLE_MIN_LENGTH,
    FIELD_ALIASES,
    MIN_CONTENT_LENGTH,
    SUPPORTED_SUFFIXES,
)


# ══════════════════════════════════════════════════════════════════
# Chuẩn hoá tên cột
# ══════════════════════════════════════════════════════════════════

def _slug(text) -> str:
    """
    Đưa tên cột về dạng so khớp được: bỏ dấu, thường hoá, gạch dưới.

    Nguồn dữ liệu người Việt xuất ra hay có header tiếng Việt CÓ DẤU
    ("Nội dung", "Số ký hiệu", "Loại văn bản"). So khớp thẳng chuỗi thì cứ đổi
    nguồn là trượt hết. Bỏ dấu trước rồi mới khớp: "Nội dung" thành `noi_dung`,
    trùng luôn với bí danh đã khai.
    """
    text = str(text or "").strip().lower()
    text = unicodedata.normalize("NFD", text)
    text = "".join(c for c in text if unicodedata.category(c) != "Mn")
    text = text.replace("đ", "d")
    text = re.sub(r"[^a-z0-9]+", "_", text)
    return text.strip("_")


def _slug_record(raw: dict) -> dict:
    """Bản sao của record với key đã slug hoá. Giữ nguyên giá trị."""
    return {_slug(k): v for k, v in raw.items()}


# Nhiều nguồn ghi ô trống thành chuỗi chứ không phải null
_EMPTY_STRINGS = {"none", "nan", "null", "n/a", "-", "chưa xác định"}


def _pick(slugged: dict, field: str) -> str:
    """Lấy giá trị đầu tiên không rỗng theo danh sách bí danh của `field`."""
    for alias in FIELD_ALIASES.get(field, []):
        value = slugged.get(_slug(alias))
        if value is None:
            continue
        text = str(value).strip()
        if text and text.lower() not in _EMPTY_STRINGS:
            return text
    return ""


# ══════════════════════════════════════════════════════════════════
# Đọc file thô thành list record dict
# ══════════════════════════════════════════════════════════════════

_HTML_RE = re.compile(r"<\s*(p|div|br|span|table|tr|td|h[1-6]|body|html)\b", re.IGNORECASE)


def _read_text(path: Path) -> str:
    """
    Đọc file text, thử lần lượt vài bảng mã.

    `utf-8-sig` đứng đầu vì Excel trên Windows xuất CSV kèm BOM; không xử lý
    BOM thì cột đầu tiên mang tên rác và mọi bí danh đều trượt.
    """
    for encoding in ("utf-8-sig", "utf-8", "cp1258", "cp1252", "latin-1"):
        try:
            return path.read_text(encoding=encoding)
        except UnicodeDecodeError:
            continue
    return path.read_bytes().decode("utf-8", errors="replace")


def _read_jsonl(path: Path) -> list[dict]:
    records = []
    for line_no, line in enumerate(_read_text(path).splitlines(), 1):
        line = line.strip()
        if not line:
            continue
        try:
            obj = json.loads(line)
        except json.JSONDecodeError as exc:
            logger.warning(f"{path.name}:{line_no} — bỏ qua dòng JSON hỏng ({exc})")
            continue
        if isinstance(obj, dict):
            records.append(obj)
    return records


def _read_json(path: Path) -> list[dict]:
    data = json.loads(_read_text(path))
    if isinstance(data, dict):
        # Bọc kiểu {"data": [...]} / {"documents": [...]} rất phổ biến
        for key in ("data", "documents", "docs", "records", "items", "results"):
            if isinstance(data.get(key), list):
                data = data[key]
                break
        else:
            data = [data]
    return [d for d in data if isinstance(d, dict)]


def _read_csv(path: Path) -> list[dict]:
    text = _read_text(path)
    try:
        dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|")
    except csv.Error:
        # Sniffer hay thua khi ô chứa văn bản luật dài có đủ loại dấu câu.
        # Đoán theo phần mở rộng thay vì bỏ cuộc.
        dialect = csv.excel_tab if path.suffix.lower() == ".tsv" else csv.excel
    # Văn bản luật dễ vượt giới hạn field mặc định (128KB) của module csv
    csv.field_size_limit(2 ** 31 - 1)
    return [dict(row) for row in csv.DictReader(text.splitlines(), dialect=dialect)]


def _read_xlsx(path: Path) -> list[dict]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ImportError(
            f"Cần `openpyxl` để đọc {path.name}. Cài: pip install openpyxl"
        ) from exc

    workbook = load_workbook(path, read_only=True, data_only=True)
    records = []
    for sheet in workbook.worksheets:
        rows = sheet.iter_rows(values_only=True)
        try:
            header = [str(c) if c is not None else "" for c in next(rows)]
        except StopIteration:
            continue
        for row in rows:
            if all(c is None or str(c).strip() == "" for c in row):
                continue
            records.append({h: v for h, v in zip(header, row) if h})
    workbook.close()
    return records


def _read_docx(path: Path) -> list[dict]:
    """
    Một file .docx là MỘT văn bản luật.

    Lấy cả text trong bảng: biểu mẫu, phụ lục, bảng mức bồi thường trong văn
    bản luật Việt Nam thường nằm trong table, bỏ qua là mất đúng phần nội dung
    đáng hỏi nhất.
    """
    try:
        import docx
    except ImportError as exc:
        raise ImportError(
            f"Cần `python-docx` để đọc {path.name}. Cài: pip install python-docx"
        ) from exc

    document = docx.Document(str(path))
    parts = [p.text for p in document.paragraphs if p.text and p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text and c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))

    return [{"title": path.stem, "content": "\n".join(parts)}]


def _read_plain(path: Path) -> list[dict]:
    """Một file .txt/.md/.html là MỘT văn bản luật; tên file làm tiêu đề."""
    return [{"title": path.stem, "content": _read_text(path)}]


def _read_parquet(path: Path) -> list[dict]:
    try:
        import pyarrow.parquet as pq
    except ImportError as exc:
        raise ImportError(
            f"Cần `pyarrow` để đọc {path.name}. Cài: pip install pyarrow"
        ) from exc
    return pq.read_table(path).to_pylist()


_READERS = {
    ".jsonl": _read_jsonl, ".ndjson": _read_jsonl,
    ".json": _read_json,
    ".csv": _read_csv, ".tsv": _read_csv,
    ".xlsx": _read_xlsx, ".xlsm": _read_xlsx,
    ".docx": _read_docx,
    ".txt": _read_plain, ".md": _read_plain,
    ".html": _read_plain, ".htm": _read_plain,
    ".parquet": _read_parquet,
}

# Hai nơi khai định dạng phải khớp nhau, nếu không thông báo lỗi cho người
# dùng sẽ liệt kê một đằng còn code đọc được một nẻo.
assert set(_READERS) == set(SUPPORTED_SUFFIXES), (
    "SUPPORTED_SUFFIXES trong config/qa_settings.py đã lệch với _READERS. Sửa cả hai."
)


# ══════════════════════════════════════════════════════════════════
# Chuẩn hoá record thành doc dict
# ══════════════════════════════════════════════════════════════════

def _to_plain_text(value) -> str:
    text = "" if value is None else str(value)
    if not _HTML_RE.search(text):
        return text.strip()

    # Import muộn: `src.data_loader` kéo theo `datasets` và `config.settings`.
    # Đường chạy cục bộ không cần hai thứ đó và không nên chết vì thiếu chúng.
    try:
        from src.data_loader import clean_html
        return clean_html(text)
    except Exception:
        from bs4 import BeautifulSoup
        return BeautifulSoup(text, "html.parser").get_text(separator="\n").strip()


# Các trường metadata bê nguyên nếu nguồn có
_METADATA_FIELDS = (
    "title", "so_hieu", "loai_van_ban", "nganh", "linh_vuc",
    "co_quan_ban_hanh", "ngay_ban_hanh", "tinh_trang_hieu_luc",
)


def normalize_record(raw: dict, source_file: str = "", index: int = 0) -> dict | None:
    """
    Ép một record thô bất kỳ về doc dict chuẩn hoá.

    Trả None khi record không có nội dung — đó là chuyện bình thường (dòng
    trống trong CSV, sheet phụ trong Excel), không phải lỗi cần báo.
    """
    slugged = _slug_record(raw)
    content = _to_plain_text(_pick(slugged, "content"))
    if not content:
        return None

    metadata = {field: _pick(slugged, field) for field in _METADATA_FIELDS}
    # Giữ lại mọi cột lạ của nguồn: khi phát hiện dữ liệu bẩn còn có cái mà
    # truy, và nguồn mới có thể mang trường hữu ích mà ta chưa biết trước.
    for key, value in slugged.items():
        if key not in metadata and value not in (None, ""):
            metadata.setdefault(key, str(value).strip())

    doc_id = _pick(slugged, "id")
    if not doc_id:
        from src.utils import generate_item_id
        seed = metadata.get("so_hieu") or metadata.get("title") or content[:500]
        doc_id = generate_item_id(f"{source_file}:{index}:{seed}")

    return {
        "id": str(doc_id),
        "content": content,
        "content_length": len(content),
        "metadata": metadata,
        "source_doc_id": f"doc_{doc_id}",
        "source_file": source_file,
    }


# ══════════════════════════════════════════════════════════════════
# Điểm vào
# ══════════════════════════════════════════════════════════════════

def iter_source_files(source) -> list[Path]:
    """Liệt kê file đọc được trong `source` (file đơn hoặc thư mục, đệ quy)."""
    source = Path(source)
    if source.is_file():
        return [source] if source.suffix.lower() in _READERS else []
    if not source.is_dir():
        return []
    return sorted(
        p for p in source.rglob("*")
        # "~$..." là file khoá tạm của Word/Excel, mở ra là hỏng
        if p.is_file() and p.suffix.lower() in _READERS and not p.name.startswith("~$")
    )


def has_local_corpus(source) -> bool:
    """Có file nào đọc được ở `source` không? Quyết định chọn đường nạp."""
    return bool(source) and bool(iter_source_files(source))


def load_local_corpus(source, min_content_length: int | None = None) -> list[dict]:
    """
    Nạp toàn bộ văn bản từ `source` về danh sách doc dict chuẩn hoá.

    Args:
        source: file hoặc thư mục. Thư mục thì quét đệ quy.
        min_content_length: ngưỡng độ dài; None thì lấy `MIN_CONTENT_LENGTH`.

    Trùng `source_doc_id` bị loại (giữ bản đầu): nguồn tải tay rất hay có file
    lặp, mà trùng document giữa train và test là data leakage kinh điển.
    """
    threshold = MIN_CONTENT_LENGTH if min_content_length is None else min_content_length
    files = iter_source_files(source)
    if not files:
        logger.warning(f"Không tìm thấy file đọc được nào trong: {source}")
        return []

    logger.info(f"Nạp corpus cục bộ từ {source} — {len(files)} file")

    documents: list[dict] = []
    seen: set[str] = set()
    skipped_short = 0
    skipped_dup = 0

    for path in files:
        reader = _READERS[path.suffix.lower()]
        try:
            raw_records = reader(path)
        except Exception as exc:
            # Một file hỏng không được phép làm hỏng cả mẻ nạp
            logger.error(f"Bỏ qua {path.name}: {type(exc).__name__}: {exc}")
            continue

        n_before = len(documents)
        for index, raw in enumerate(raw_records):
            doc = normalize_record(raw, source_file=path.name, index=index)
            if doc is None:
                continue
            if doc["content_length"] < threshold:
                skipped_short += 1
                continue
            if doc["source_doc_id"] in seen:
                skipped_dup += 1
                continue
            seen.add(doc["source_doc_id"])
            documents.append(doc)

        logger.debug(f"  {path.name}: {len(documents) - n_before} document")

    logger.info(
        f"Corpus cục bộ: {len(documents)} document "
        f"({skipped_short} bị loại vì content < {threshold} ký tự, "
        f"{skipped_dup} trùng source_doc_id)"
    )
    return documents


# ══════════════════════════════════════════════════════════════════
# Cắt theo Điều (tuỳ chọn — xem QA_SPLIT_BY_ARTICLE)
# ══════════════════════════════════════════════════════════════════

# "Điều 5.", "Điều 12:", "Điều 468a." — đầu dòng, có số, có dấu ngắt.
# Cố ý KHÔNG bắt "điều" viết thường giữa câu ("theo điều 5 nói trên").
_ARTICLE_RE = re.compile(r"^[ \t]*Điều\s+(\d+[a-zđ]?)\s*[.:．]?\s", re.MULTILINE)


def split_into_articles(content: str, min_length: int | None = None) -> list[tuple[str, str]]:
    """
    Cắt nội dung văn bản thành từng Điều.

    Returns:
        list[(số điều, nội dung điều)]. Rỗng khi văn bản không đánh số theo
        Điều (Thông tư đánh `I.`/`1.`, văn bản sửa đổi...) — caller giữ nguyên
        cả văn bản trong trường hợp đó.
    """
    threshold = ARTICLE_MIN_LENGTH if min_length is None else min_length
    matches = list(_ARTICLE_RE.finditer(content or ""))
    if len(matches) < 2:
        return []

    articles: list[tuple[str, str]] = []
    for i, match in enumerate(matches):
        end = matches[i + 1].start() if i + 1 < len(matches) else len(content)
        body = content[match.start():end].strip()
        if len(body) >= threshold:
            articles.append((match.group(1), body))
    return articles


def expand_documents_by_article(documents: list[dict]) -> list[dict]:
    """
    Biến mỗi văn bản thành nhiều "document con", mỗi Điều một cái.

    Vì sao cần: phạm vi dân sự là ÍT VĂN BẢN NHƯNG RẤT DÀI (riêng Bộ luật Dân
    sự 2015 có 689 Điều). Lấy 1 cặp câu hỏi trên mỗi văn bản thì cả bộ luật
    lớn nhất cũng chỉ đóng góp đúng 1 cặp, mà 2.000 ký tự đầu của nó thì toàn
    phần "Căn cứ..." — nội dung đáng hỏi nằm ở giữa và không bao giờ tới lượt.

    `source_doc_id` chuyển thành `<doc>#dieu-<n>` và ĐÓ MỚI là khoá group
    split. Cách đọc: hai câu của cùng một Điều không bao giờ tách sang hai
    split khác nhau (đây là leakage thật cần chặn), còn các Điều khác nhau
    trong cùng một bộ luật thì được phép nằm khác split — nếu không thì cả bộ
    luật rơi trọn vào một split và không chia nổi train/val/test. Phải ghi rõ
    cách đọc này ở phần Limitations của báo cáo.
    """
    expanded: list[dict] = []
    n_split = 0

    for doc in documents:
        articles = split_into_articles(doc.get("content", ""))
        if not articles:
            expanded.append(doc)
            continue

        n_split += 1
        for number, body in articles:
            child = dict(doc)
            child["content"] = body
            child["content_length"] = len(body)
            child["source_doc_id"] = f"{doc['source_doc_id']}#dieu-{number}"
            child["id"] = f"{doc['id']}#dieu-{number}"
            child["parent_doc_id"] = doc["source_doc_id"]
            child["article_number"] = number
            child["metadata"] = dict(doc.get("metadata") or {})
            child["metadata"]["dieu"] = number
            expanded.append(child)

    logger.info(
        f"Cắt theo Điều: {n_split}/{len(documents)} văn bản cắt được "
        f"→ {len(expanded)} document con"
    )
    return expanded
